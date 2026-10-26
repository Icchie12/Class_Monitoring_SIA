import os
import re
import json
import sqlite3
import base64
from email.message import EmailMessage
from datetime import datetime, timezone
from zoneinfo import ZoneInfo
from io import BytesIO
from urllib.parse import urlparse

import requests
from flask import Flask, jsonify, request, render_template, send_file, redirect, session
from flask_socketio import SocketIO, emit
from openpyxl import Workbook
from openpyxl.styles import Font
from google_auth_oauthlib.flow import Flow
from google.oauth2.credentials import Credentials
from google.auth.transport.requests import Request as GoogleRequest
from googleapiclient.discovery import build

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "monitor.db")
GOOGLE_TOKEN_PATH = os.path.join(BASE_DIR, "google_token.json")
GOOGLE_CLIENT_SECRET = os.path.join(BASE_DIR, "credentials.json")

app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", "school-prototype-secret")
socketio = SocketIO(app, cors_allowed_origins="*")

INACTIVE_SECONDS = 300
POLL_SECONDS = 5
# Local school-project prototype only. This allows the OAuth callback to use
# http://127.0.0.1 during local development. Do not use this setting in production.
os.environ.setdefault("OAUTHLIB_INSECURE_TRANSPORT", "1")

SCOPES = [
    "https://www.googleapis.com/auth/meetings.space.readonly",
    "https://www.googleapis.com/auth/gmail.send",
]


def now_iso():
    return datetime.now(timezone.utc).isoformat()

def manila_iso(value):
    """Convert an ISO timestamp to Philippine time and keep +08:00 visible in Excel."""
    if not value:
        return ""
    try:
        dt = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        dt = dt.astimezone(ZoneInfo("Asia/Manila"))
        return dt.strftime("%Y-%m-%d %I:%M:%S %p +08:00")
    except Exception:
        return str(value)


def db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = db()
    conn.executescript("""
    CREATE TABLE IF NOT EXISTS sessions (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        class_name TEXT NOT NULL,
        teacher_email TEXT NOT NULL,
        meeting_url TEXT,
        meeting_code TEXT,
        started_at TEXT NOT NULL,
        ended_at TEXT,
        conference_record TEXT
    );

    CREATE TABLE IF NOT EXISTS students (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        session_id INTEGER NOT NULL,
        student_name TEXT NOT NULL,
        extension_id TEXT DEFAULT '',
        google_user_id TEXT DEFAULT '',
        time_in TEXT,
        time_out TEXT,
        last_activity TEXT,
        status TEXT DEFAULT 'WAITING',
        active_seconds INTEGER DEFAULT 0,
        inactive_seconds INTEGER DEFAULT 0,
        tab_changes INTEGER DEFAULT 0,
        camera_events INTEGER DEFAULT 0,
        microphone_events INTEGER DEFAULT 0,
        chat_events INTEGER DEFAULT 0,
        FOREIGN KEY(session_id) REFERENCES sessions(id)
    );

    CREATE TABLE IF NOT EXISTS events (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        student_id INTEGER NOT NULL,
        event_type TEXT NOT NULL,
        event_time TEXT NOT NULL,
        details TEXT,
        FOREIGN KEY(student_id) REFERENCES students(id)
    );
    """)

    # Upgrade old prototype databases safely.
    existing = {r[1] for r in conn.execute("PRAGMA table_info(sessions)").fetchall()}
    for col, definition in [
        ("meeting_code", "TEXT"),
        ("conference_record", "TEXT"),
    ]:
        if col not in existing:
            conn.execute(f"ALTER TABLE sessions ADD COLUMN {col} {definition}")

    existing = {r[1] for r in conn.execute("PRAGMA table_info(students)").fetchall()}
    for col, definition in [
        ("google_user_id", "TEXT DEFAULT ''"),
    ]:
        if col not in existing:
            conn.execute(f"ALTER TABLE students ADD COLUMN {col} {definition}")

    conn.commit()
    conn.close()


def extract_meeting_code(meeting_url):
    """Extract abc-defg-hij from a normal Google Meet URL."""
    if not meeting_url:
        return ""
    value = meeting_url.strip()
    match = re.search(r"meet\.google\.com/([a-z0-9-]+)", value, re.I)
    if match:
        return match.group(1).split("?")[0].split("#")[0]
    return ""


def load_google_credentials():
    if not os.path.exists(GOOGLE_TOKEN_PATH):
        return None
    try:
        with open(GOOGLE_TOKEN_PATH, "r", encoding="utf-8") as f:
            data = json.load(f)
        creds = Credentials.from_authorized_user_info(data, SCOPES)
        if creds.expired and creds.refresh_token:
            creds.refresh(GoogleRequest())
            save_google_credentials(creds)
        return creds if creds.valid else None
    except Exception as exc:
        print("Google token load error:", exc)
        return None


def save_google_credentials(creds):
    with open(GOOGLE_TOKEN_PATH, "w", encoding="utf-8") as f:
        f.write(creds.to_json())


def google_auth_configured():
    return os.path.exists(GOOGLE_CLIENT_SECRET)


def google_api_get(path, creds, params=None):
    headers = {"Authorization": f"Bearer {creds.token}"}
    response = requests.get(
        "https://meet.googleapis.com/v2/" + path.lstrip("/"),
        headers=headers,
        params=params or {},
        timeout=15,
    )
    if response.status_code == 401 and creds.refresh_token:
        creds.refresh(GoogleRequest())
        save_google_credentials(creds)
        headers["Authorization"] = f"Bearer {creds.token}"
        response = requests.get(
            "https://meet.googleapis.com/v2/" + path.lstrip("/"),
            headers=headers,
            params=params or {},
            timeout=15,
        )
    response.raise_for_status()
    return response.json()


def list_active_conference(meeting_code, creds):
    data = google_api_get(
        "conferenceRecords",
        creds,
        params={
            "filter": f'space.meeting_code = "{meeting_code}" AND end_time IS NULL',
            "pageSize": 10,
        },
    )
    records = data.get("conferenceRecords", [])
    if not records:
        return None
    records.sort(key=lambda x: x.get("startTime", ""), reverse=True)
    return records[0]


def list_participants(conference_name, creds):
    participants = []
    page_token = None
    while True:
        params = {"pageSize": 250, "filter": "latest_end_time IS NULL"}
        if page_token:
            params["pageToken"] = page_token
        path = f"{conference_name}/participants".replace("//", "/")
        data = google_api_get(path, creds, params=params)
        participants.extend(data.get("participants", []))
        page_token = data.get("nextPageToken")
        if not page_token:
            break
    return participants


def participant_display_name(participant):
    user = participant.get("signedinUser")
    if user:
        return user.get("displayName", "Google user"), user.get("user", "")
    anonymous = participant.get("anonymousUser")
    if anonymous:
        return anonymous.get("displayName", "Anonymous participant"), ""
    phone = participant.get("phoneUser")
    if phone:
        return phone.get("displayName", "Phone participant"), ""
    return "Unknown participant", ""


def sync_google_participants(session_id):
    """Synchronize active Meet participants into the dashboard."""
    creds = load_google_credentials()
    if not creds:
        return {"ok": False, "error": "Google account is not connected."}

    conn = db()
    session_row = conn.execute("SELECT * FROM sessions WHERE id=?", (session_id,)).fetchone()
    if not session_row:
        conn.close()
        return {"ok": False, "error": "Session not found."}

    meeting_code = session_row["meeting_code"] or extract_meeting_code(session_row["meeting_url"])
    if not meeting_code:
        conn.close()
        return {"ok": False, "error": "Invalid Google Meet URL."}

    try:
        conference = list_active_conference(meeting_code, creds)
        if not conference:
            conn.close()
            return {"ok": True, "found": False, "participants": 0}

        conference_name = conference["name"]
        participants = list_participants(conference_name, creds)
        timestamp = now_iso()

        conn.execute(
            "UPDATE sessions SET conference_record=? WHERE id=?",
            (conference_name, session_id),
        )

        seen_ids = set()
        for p in participants:
            name, google_user_id = participant_display_name(p)
            if not name:
                continue
            seen_ids.add(google_user_id or f"name:{name.lower()}")
            start_time = p.get("earliestStartTime") or timestamp

            row = None
            if google_user_id:
                row = conn.execute(
                    "SELECT * FROM students WHERE session_id=? AND google_user_id=?",
                    (session_id, google_user_id),
                ).fetchone()
            if not row:
                row = conn.execute(
                    "SELECT * FROM students WHERE session_id=? AND lower(student_name)=lower(?)",
                    (session_id, name),
                ).fetchone()

            if row:
                conn.execute("""
                    UPDATE students
                    SET student_name=?, google_user_id=?,
                        time_in=COALESCE(time_in, ?),
                        time_out=NULL,
                        status='ACTIVE',
                        last_activity=COALESCE(last_activity, ?)
                    WHERE id=?
                """, (name, google_user_id, start_time, timestamp, row["id"]))
                student_id = row["id"]
            else:
                cur = conn.execute("""
                    INSERT INTO students(
                        session_id, student_name, extension_id, google_user_id,
                        time_in, last_activity, status
                    ) VALUES (?, ?, '', ?, ?, ?, 'ACTIVE')
                """, (session_id, name, google_user_id, start_time, timestamp))
                student_id = cur.lastrowid
                conn.execute("""
                    INSERT INTO events(student_id,event_type,event_time,details)
                    VALUES (?, 'MEET_JOIN', ?, 'Participant detected by Google Meet API')
                """, (student_id, timestamp))

            conn.execute("""
                UPDATE students SET status='ACTIVE', time_out=NULL WHERE id=?
            """, (student_id,))

        # Mark students that disappeared from the active participant list as LEFT.
        current_rows = conn.execute(
            "SELECT * FROM students WHERE session_id=? AND status!='LEFT'",
            (session_id,),
        ).fetchall()
        for st in current_rows:
            key = st["google_user_id"] or f"name:{st['student_name'].lower()}"
            if key not in seen_ids and st["google_user_id"]:
                conn.execute("""
                    UPDATE students
                    SET status='LEFT', time_out=COALESCE(time_out, ?)
                    WHERE id=?
                """, (timestamp, st["id"]))
                conn.execute("""
                    INSERT INTO events(student_id,event_type,event_time,details)
                    VALUES (?, 'MEET_LEFT', ?, 'Participant no longer appears in active Meet sessions')
                """, (st["id"], timestamp))

        conn.commit()
        conn.close()

        # Broadcast the whole session so the teacher UI refreshes immediately.
        socketio.emit("participants_sync", {"session_id": session_id})
        return {"ok": True, "found": True, "participants": len(participants)}
    except requests.HTTPError as exc:
        conn.close()
        body = exc.response.text[:500] if exc.response is not None else str(exc)
        print("Google Meet API error:", body)
        return {"ok": False, "error": f"Google Meet API error: {body}"}
    except Exception as exc:
        conn.close()
        print("Participant sync error:", exc)
        return {"ok": False, "error": str(exc)}


def emit_student_update(student_id):
    conn = db()
    row = conn.execute("""
        SELECT st.*, s.class_name
        FROM students st JOIN sessions s ON st.session_id=s.id
        WHERE st.id=?
    """, (student_id,)).fetchone()
    conn.close()
    if row:
        socketio.emit("student_update", dict(row))


@app.route("/")
def index():
    return render_template("teacher.html", google_connected=bool(load_google_credentials()))


@app.route("/student")
def student_page():
    return render_template("student.html")


@app.route("/auth/google")
def google_auth():
    if not google_auth_configured():
        return "<h2>Google OAuth is not configured.</h2><p>Place credentials.json from Google Cloud Console in the project folder.</p>", 500

    flow = Flow.from_client_secrets_file(
        GOOGLE_CLIENT_SECRET,
        scopes=SCOPES,
        redirect_uri="http://127.0.0.1:5000/oauth2callback",
    )
    authorization_url, state = flow.authorization_url(
        access_type="offline",
        include_granted_scopes="true",
        prompt="consent",
    )
    session["oauth_state"] = state
    return redirect(authorization_url)


@app.route("/oauth2callback")
def oauth2callback():
    if not google_auth_configured():
        return "Google OAuth is not configured.", 500

    flow = Flow.from_client_secrets_file(
        GOOGLE_CLIENT_SECRET,
        scopes=SCOPES,
        state=session.get("oauth_state"),
        redirect_uri="http://127.0.0.1:5000/oauth2callback",
    )
    flow.fetch_token(authorization_response=request.url)
    save_google_credentials(flow.credentials)
    return redirect("/")


@app.route("/api/google/status")
def google_status():
    creds = load_google_credentials()
    return jsonify({
        "configured": google_auth_configured(),
        "connected": bool(creds),
    })


@app.route("/api/session", methods=["POST"])
def create_session():
    data = request.json or {}
    class_name = data.get("class_name", "").strip()
    teacher_email = data.get("teacher_email", "").strip()
    meeting_url = data.get("meeting_url", "").strip()
    meeting_code = extract_meeting_code(meeting_url)

    if not class_name or not teacher_email or not meeting_code:
        return jsonify({"error": "Enter class name, teacher email, and a valid Google Meet URL."}), 400

    if not load_google_credentials():
        return jsonify({"error": "Connect the teacher's Google account first."}), 401

    conn = db()
    cur = conn.execute("""
        INSERT INTO sessions(class_name, teacher_email, meeting_url, meeting_code, started_at)
        VALUES (?, ?, ?, ?, ?)
    """, (class_name, teacher_email, meeting_url, meeting_code, now_iso()))
    conn.commit()
    session_id = cur.lastrowid
    conn.close()

    # Do an immediate sync so participants already inside Meet appear right away.
    result = sync_google_participants(session_id)
    return jsonify({"session_id": session_id, "sync": result})


@app.route("/api/session/<int:session_id>", methods=["GET"])
def get_session(session_id):
    conn = db()
    session_row = conn.execute("SELECT * FROM sessions WHERE id=?", (session_id,)).fetchone()
    students = conn.execute("""
        SELECT * FROM students WHERE session_id=? ORDER BY student_name
    """, (session_id,)).fetchall()
    conn.close()

    if not session_row:
        return jsonify({"error": "Session not found"}), 404

    return jsonify({
        "session": dict(session_row),
        "students": [dict(x) for x in students]
    })


@app.route("/api/session/<int:session_id>/sync", methods=["POST"])
def manual_sync(session_id):
    return jsonify(sync_google_participants(session_id))


@app.route("/api/student/join", methods=["POST"])
def student_join():
    data = request.json or {}
    session_id = data.get("session_id")
    name = data.get("student_name", "").strip()
    extension_id = data.get("extension_id", "").strip()

    if not session_id or not name:
        return jsonify({"error": "session_id and student_name are required"}), 400

    conn = db()
    session_row = conn.execute("SELECT * FROM sessions WHERE id=?", (session_id,)).fetchone()
    if not session_row:
        conn.close()
        return jsonify({"error": "Session not found"}), 404

    existing = conn.execute("""
        SELECT id FROM students
        WHERE session_id=? AND lower(student_name)=lower(?)
    """, (session_id, name)).fetchone()

    timestamp = now_iso()
    if existing:
        student_id = existing["id"]
        conn.execute("""
            UPDATE students
            SET extension_id=?, time_in=COALESCE(time_in, ?),
                last_activity=?, status='ACTIVE'
            WHERE id=?
        """, (extension_id, timestamp, timestamp, student_id))
    else:
        cur = conn.execute("""
            INSERT INTO students(
                session_id, student_name, extension_id, time_in,
                last_activity, status
            ) VALUES (?, ?, ?, ?, ?, 'ACTIVE')
        """, (session_id, name, extension_id, timestamp, timestamp))
        student_id = cur.lastrowid

    conn.execute("""
        INSERT INTO events(student_id,event_type,event_time,details)
        VALUES (?, 'STUDENT_MONITOR_START', ?, 'Student monitoring extension connected')
    """, (student_id, timestamp))
    conn.commit()
    conn.close()
    emit_student_update(student_id)
    return jsonify({"student_id": student_id, "session_id": session_id})


@app.route("/api/event", methods=["POST"])
def record_event():
    data = request.json or {}
    student_id = data.get("student_id")
    event_type = data.get("event_type")
    details = data.get("details", "")

    allowed = {
        "TAB_HIDDEN", "TAB_VISIBLE", "CAMERA_ON", "CAMERA_OFF",
        "MIC_ON", "MIC_OFF", "CHAT", "ACTIVITY", "LEFT_MEET",
        "RETURNED_MEET", "HEARTBEAT"
    }
    if not student_id or event_type not in allowed:
        return jsonify({"error": "Invalid event"}), 400

    timestamp = now_iso()
    conn = db()
    row = conn.execute("SELECT * FROM students WHERE id=?", (student_id,)).fetchone()
    if not row:
        conn.close()
        return jsonify({"error": "Student not found"}), 404

    # Activity pings update status only; they are not written to the event log.
    if event_type in ("ACTIVITY", "HEARTBEAT"):
        conn.execute("UPDATE students SET last_activity=?, status='ACTIVE' WHERE id=?", (timestamp, student_id))
        conn.commit()
        conn.close()
        emit_student_update(student_id)
        return jsonify({"ok": True, "logged": False})

    tab_changes = row["tab_changes"]
    camera_events = row["camera_events"]
    microphone_events = row["microphone_events"]
    chat_events = row["chat_events"]

    if event_type == "TAB_HIDDEN":
        tab_changes += 1
        status = "AWAY"
    elif event_type in ("TAB_VISIBLE", "RETURNED_MEET", "ACTIVITY", "CHAT", "CAMERA_ON", "MIC_ON"):
        status = "ACTIVE"
    elif event_type == "LEFT_MEET":
        status = "LEFT"
    else:
        status = row["status"]

    if event_type.startswith("CAMERA_"):
        camera_events += 1
    if event_type.startswith("MIC_"):
        microphone_events += 1
    if event_type == "CHAT":
        chat_events += 1

    conn.execute("""
        INSERT INTO events(student_id,event_type,event_time,details)
        VALUES (?, ?, ?, ?)
    """, (student_id, event_type, timestamp, details))
    conn.execute("""
        UPDATE students
        SET last_activity=?, status=?, tab_changes=?,
            camera_events=?, microphone_events=?, chat_events=?
        WHERE id=?
    """, (timestamp, status, tab_changes, camera_events, microphone_events, chat_events, student_id))
    conn.commit()
    conn.close()
    emit_student_update(student_id)
    return jsonify({"ok": True})


@app.route("/api/heartbeat", methods=["POST"])
def heartbeat():
    data = request.json or {}
    student_id = data.get("student_id")
    active = bool(data.get("active", False))
    if not student_id:
        return jsonify({"error": "student_id required"}), 400

    conn = db()
    row = conn.execute("SELECT * FROM students WHERE id=?", (student_id,)).fetchone()
    if not row:
        conn.close()
        return jsonify({"error": "Student not found"}), 404

    timestamp = now_iso()
    status = "ACTIVE" if active else "INACTIVE"
    conn.execute("""
        UPDATE students
        SET last_activity=?, status=?
        WHERE id=?
    """, (timestamp, status, student_id))
    conn.commit()
    conn.close()
    emit_student_update(student_id)
    return jsonify({"ok": True})


@app.route("/api/student/<int:student_id>/leave", methods=["POST"])
def student_leave(student_id):
    timestamp = now_iso()
    conn = db()
    conn.execute("UPDATE students SET time_out=?, status='LEFT' WHERE id=?", (timestamp, student_id))
    conn.execute("""
        INSERT INTO events(student_id,event_type,event_time,details)
        VALUES (?, 'LEAVE', ?, 'Student stopped monitoring')
    """, (student_id, timestamp))
    conn.commit()
    conn.close()
    emit_student_update(student_id)
    return jsonify({"ok": True})


@app.route("/api/session/<int:session_id>/events/<int:student_id>")
def student_events(session_id, student_id):
    conn = db()
    events = conn.execute("""
        SELECT e.* FROM events e
        JOIN students st ON e.student_id=st.id
        WHERE e.student_id=? AND st.session_id=?
        ORDER BY e.event_time ASC
    """, (student_id, session_id)).fetchall()
    conn.close()
    return jsonify([dict(x) for x in events])


def make_report(session_id):
    conn = db()
    session_row = conn.execute("SELECT * FROM sessions WHERE id=?", (session_id,)).fetchone()
    students = conn.execute("SELECT * FROM students WHERE session_id=? ORDER BY student_name", (session_id,)).fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Summary"
    headers = [
        "Student", "Time In", "Time Out", "Status",
        "Active Time (sec)", "Inactive Time (sec)",
        "Tab Changes", "Camera Events", "Microphone Events", "Chat Events"
    ]
    ws.append(headers)
    for st in students:
        ws.append([
            st["student_name"], manila_iso(st["time_in"]), manila_iso(st["time_out"]), st["status"],
            st["active_seconds"], st["inactive_seconds"], st["tab_changes"],
            st["camera_events"], st["microphone_events"], st["chat_events"]
        ])

    ws2 = wb.create_sheet("Activity Log")
    ws2.append(["Student", "Event", "Timestamp", "Details"])
    events = conn.execute("""
        SELECT st.student_name, e.event_type, e.event_time, e.details
        FROM events e JOIN students st ON e.student_id=st.id
        WHERE st.session_id=? ORDER BY e.event_time
    """, (session_id,)).fetchall()
    for e in events:
        ws2.append([e["student_name"], e["event_type"], manila_iso(e["event_time"]), e["details"]])

    ws3 = wb.create_sheet("Class Summary")
    ws3.append(["Class", session_row["class_name"]])
    ws3.append(["Teacher Email", session_row["teacher_email"]])
    ws3.append(["Meet URL", session_row["meeting_url"]])
    ws3.append(["Meet Code", session_row["meeting_code"]])
    ws3.append(["Started", manila_iso(session_row["started_at"])])
    ws3.append(["Ended", manila_iso(session_row["ended_at"])])
    ws3.append(["Students Detected", len(students)])

    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column in sheet.columns:
            max_len = max(len(str(c.value or "")) for c in column)
            sheet.column_dimensions[column[0].column_letter].width = min(max(max_len + 2, 12), 45)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    conn.close()
    return output


@app.route("/api/session/<int:session_id>/report")
def download_report(session_id):
    output = make_report(session_id)
    return send_file(
        output,
        as_attachment=True,
        download_name=f"class_report_{session_id}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/api/session/<int:session_id>/end", methods=["POST"])
def end_session(session_id):
    timestamp = now_iso()
    conn = db()
    conn.execute("UPDATE sessions SET ended_at=? WHERE id=?", (timestamp, session_id))
    conn.execute("""
        UPDATE students SET status='LEFT', time_out=COALESCE(time_out, ?)
        WHERE session_id=? AND status != 'LEFT'
    """, (timestamp, session_id))
    conn.commit()
    session_row = conn.execute("SELECT * FROM sessions WHERE id=?", (session_id,)).fetchone()
    conn.close()
    email_sent = email_report(session_id, session_row["teacher_email"])
    socketio.emit("session_ended", {"session_id": session_id})
    return jsonify({"ok": True, "email_sent": email_sent})


def email_report(session_id, recipient):
    """Send the generated Excel report through the teacher's connected Gmail account."""
    creds = load_google_credentials()
    if not creds or not creds.valid:
        print("Email failed: Google account is not connected.")
        return False

    report = make_report(session_id)
    msg = EmailMessage()
    msg["Subject"] = f"Online Class Monitoring Report - Session {session_id}"
    msg["To"] = recipient
    msg.set_content(
        "Attached is the generated online class monitoring report for the completed class session."
    )
    msg.add_attachment(
        report.read(),
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"class_report_{session_id}.xlsx",
    )

    try:
        gmail = build("gmail", "v1", credentials=creds, cache_discovery=False)
        raw = base64.urlsafe_b64encode(msg.as_bytes()).decode("utf-8")
        gmail.users().messages().send(
            userId="me",
            body={"raw": raw}
        ).execute()
        print(f"Report email sent to {recipient}.")
        return True
    except Exception as exc:
        print("Email failed:", exc)
        return False


@socketio.on("connect")
def connected():
    emit("connected", {"ok": True})


@socketio.on("disconnect")
def disconnected():
    pass


if __name__ == "__main__":
    init_db()
    socketio.run(app, host="0.0.0.0", port=5000, debug=True)
